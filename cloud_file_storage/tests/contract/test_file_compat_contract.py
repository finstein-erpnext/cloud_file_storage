"""C1-C19 — the 19 invariants of `docs/research/content-consumers-contract.md` §5.

Gate F1: these run on **every** supported frappe ref with zero silent skips, which is why
they use the in-memory object store rather than a live bucket. `tests/test_minio_integration.py`
runs the same scenarios against real S3.

Scope note, stated rather than hidden: C1-C9 and C11-C16 are P2's own exit criteria
(PLAN §B). C10, C17, C18 and C19 are asserted here at the level the runtime core can
honestly assert them — the permission gate that P2 itself ships, path-write safety, URL
stability for business-field pointers, and the absence of interference with app-source
fixture reads. The full private-serving chain (C10) and the `get_file_path` patches (C17's
ecosystem half) belong to P3/P4 and their tests arrive with them.
"""

import csv
import gzip
import io
import os
import zipfile

import frappe
from frappe.core.doctype.file.utils import get_content_hash

from cloud_file_storage.cache import materialize, writeback
from cloud_file_storage.storage import objects
from cloud_file_storage.storage.exceptions import (
	CloudObjectNotFound,
	CloudStoragePermissionError,
	CloudStorageTransportError,
)
from cloud_file_storage.tests.markers import refusal_guard
from cloud_file_storage.tests.utils import CSO_DOCTYPE, CloudStorageTestCase, set_mode

UTF8_TEXT = "invoice total ₹1,20,000 — ok\n"

#: C10 — the only modules allowed to mint a presigned URL, each with the gate it runs first
#: and the suite that proves it:
#:   api/compat.py       — core's any-one-readable gate over the rows sharing a legacy key
#:                         (`tests/test_compat.py::TestLegacyGenerateFile`)
#:   serving/private.py  — `find_file_by_url` → `is_downloadable` → one access log
#:                         (`tests/test_serving_private.py`)
#:   serving/public.py   — `is_private = 0` and a DB-verified public File row (A19)
#:                         (`tests/test_serving_public.py`)
#:   backup/restore.py   — System Manager, the key scoped to this site's backup prefix AND
#:                         to a verified Backup Log manifest entry, and an audit row
#:                         committed BEFORE the mint
#:                         (`tests/test_backup_permissions.py`,
#:                          `tests/test_backup_restore.py::TestDownloadUrlIsScopedAndAudited`)
#: Adding a fifth is a security decision, not a refactor: it needs its own gate test here.
#:
#: This list was a valid proxy only while `engine.presign_get` was the sole wrapper: "no
#: module mentions presign_get" and "no module mints a URL" were then the same property.
#: `backup/restore.py` called boto3's `generate_presigned_url` directly, the equivalence
#: broke, and this contract stayed green through a whole phase (H-1). The matcher below now
#: keys on the `generate_presigned` PREFIX rather than one spelling.
SANCTIONED_PRESIGN_CALLERS = [
	"api/compat.py",
	"backup/restore.py",
	"serving/private.py",
	"serving/public.py",
]


def _is_minting_name(name: str) -> bool:
	"""Whether an identifier names a presigned-URL mint.

	One predicate, used by all three syntactic branches of the C10 walk, so the three cannot
	drift apart again. Prefix-matched for the boto3 family (`generate_presigned_url`,
	`generate_presigned_post`, any future sibling); exact for the app's own wrapper.

	Still a proxy: a mint reached under an alias (`from x import y as z`) or through a
	dynamically resolved attribute would evade it. A fully precise check needs to know that a
	receiver is a boto3 client, which is type inference rather than a syntactic walk. The
	achievable property is SYMMETRY — every name this knows about, matched in every form the
	walk understands — and that is what this gives.
	"""
	return name == "presign_get" or name.startswith("generate_presigned")


def patches_get_file_json(source: str) -> bool:
	"""Whether this module patches, rebinds or imports `frappe.get_file_json` (C19).

	Prose is excluded deliberately: a docstring or comment naming the function cannot
	patch it, and treating a mention as an offence would push accurate documentation out of
	the codebase to keep a gate green. Every executable form is still caught — string
	arguments (`patch("frappe.get_file_json")`), attribute rebinding, and `from` imports —
	which `test_C19_the_patch_detector_still_catches_every_patching_form` proves.
	"""
	import ast

	tree = ast.parse(source)

	docstrings = set()
	for node in ast.walk(tree):
		if isinstance(node, ast.Module | ast.ClassDef | ast.FunctionDef | ast.AsyncFunctionDef):
			body = getattr(node, "body", None)
			if (
				body
				and isinstance(body[0], ast.Expr)
				and isinstance(body[0].value, ast.Constant)
				and isinstance(body[0].value.value, str)
			):
				docstrings.add(id(body[0].value))

	for node in ast.walk(tree):
		# a patch by string name, e.g. patch("frappe.get_file_json")
		if isinstance(node, ast.Constant) and isinstance(node.value, str):
			if "get_file_json" in node.value and id(node) not in docstrings:
				return True
		# a direct rebinding, e.g. `frappe.get_file_json = ours` — an assignment whose
		# target is the attribute, which a Constant walk cannot see
		if isinstance(node, ast.Assign):
			for target in node.targets:
				if isinstance(target, ast.Attribute) and target.attr == "get_file_json":
					return True
		# `from frappe import get_file_json`
		if isinstance(node, ast.ImportFrom):
			if any(alias.name == "get_file_json" for alias in node.names):
				return True

	return False


def _png_bytes() -> bytes:
	from PIL import Image

	buffer = io.BytesIO()
	Image.new("RGB", (4, 4), (10, 20, 30)).save(buffer, format="PNG")
	return buffer.getvalue()


def _xlsx_bytes() -> bytes:
	from openpyxl import Workbook

	workbook = Workbook()
	workbook.active["A1"] = "hello"
	buffer = io.BytesIO()
	workbook.save(buffer)
	return buffer.getvalue()


def _zip_bytes() -> bytes:
	buffer = io.BytesIO()
	with zipfile.ZipFile(buffer, "w") as archive:
		archive.writestr("inner.txt", "inner payload")
	return buffer.getvalue()


def _gz_bytes() -> bytes:
	return gzip.compress(b'{"gstr": "payload"}')


class ContractTestCase(CloudStorageTestCase):
	"""S3_ONLY by default: every invariant is asserted with zero bytes on local disk."""

	MODE = "S3_ONLY"


class TestC1C2GetContentSignature(ContractTestCase):
	def test_C1_get_content_accepts_the_encodings_keyword(self):
		"""C1 — `gst_return_log.py:188` calls `file.get_content(encodings=[])`."""
		import inspect

		doc = self.make_file(file_name="c1.txt", content=UTF8_TEXT)
		signature = inspect.signature(type(doc).get_content)

		self.assertIn("encodings", signature.parameters)
		self.assertIsNone(signature.parameters["encodings"].default)
		# and it is callable with the keyword, not merely declared
		self.assertIsNotNone(doc.get_content(encodings=[]))

	def test_C2_encodings_empty_list_returns_raw_bytes_undecoded(self):
		"""C2 — the payload is gzip and goes straight to `as_raw()`."""
		payload = _gz_bytes()
		doc = self.make_file(file_name="c2.json.gz", content=payload)

		content = frappe.get_doc("File", doc.name).get_content(encodings=[])

		self.assertIsInstance(content, bytes)
		self.assertEqual(content, payload)
		self.assertEqual(gzip.decompress(content), b'{"gstr": "payload"}')

	def test_C2_utf8_text_is_still_bytes_when_encodings_is_empty(self):
		doc = self.make_file(file_name="c2-text.txt", content=UTF8_TEXT)
		content = frappe.get_doc("File", doc.name).get_content(encodings=[])
		self.assertIsInstance(content, bytes)
		self.assertEqual(content, UTF8_TEXT.encode())


class TestC3C4DecodePolicy(ContractTestCase):
	def test_C3_binary_content_comes_back_as_bytes_with_no_arguments(self):
		"""C3 — `pdf.py:291` b64-encodes, `gzip.decompress` needs bytes-like."""
		cases = {
			"c3.png": _png_bytes(),
			"c3.xlsx": _xlsx_bytes(),
			"c3.zip": _zip_bytes(),
			"c3.json.gz": _gz_bytes(),
			"c3.pdf": b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\ntrailer\n",
		}
		for file_name, payload in cases.items():
			with self.subTest(file_name=file_name):
				doc = self.make_file(file_name=file_name, content=payload)
				content = frappe.get_doc("File", doc.name).get_content()
				self.assertIsInstance(content, bytes)
				self.assertEqual(content, payload)

	def test_C4_utf8_decodable_content_comes_back_as_str(self):
		"""C4 — `stock_ledger.py:321-323` re-encodes when it sees `str`."""
		doc = self.make_file(file_name="c4.txt", content=UTF8_TEXT)
		content = frappe.get_doc("File", doc.name).get_content()

		self.assertIsInstance(content, str)
		self.assertEqual(content, UTF8_TEXT)

	def test_C4_round_trip_is_byte_identical_either_way(self):
		for file_name, payload in (("c4-rt.txt", UTF8_TEXT.encode()), ("c4-rt.bin", _png_bytes())):
			with self.subTest(file_name=file_name):
				doc = self.make_file(file_name=file_name, content=payload)
				content = frappe.get_doc("File", doc.name).get_content()
				as_bytes = content.encode() if isinstance(content, str) else content
				self.assertEqual(as_bytes, payload)

	def test_C4_an_explicit_encoding_list_is_tried_in_order(self):
		doc = self.make_file(file_name="c4-encodings.txt", content=UTF8_TEXT)
		reloaded = frappe.get_doc("File", doc.name)

		self.assertEqual(reloaded.get_content(encodings=["utf-8"]), UTF8_TEXT)
		# nothing applies -> bytes, never an exception
		binary = self.make_file(file_name="c4-binary.bin", content=b"\xff\xfe\x00\x01")
		self.assertIsInstance(frappe.get_doc("File", binary.name).get_content(encodings=["utf-8"]), bytes)


class TestC5ErrorTyping(ContractTestCase):
	def test_C5_content_is_readable_with_nothing_on_local_disk(self):
		doc = self.make_file(file_name="c5.txt", content="cloud resident")
		self.assertFalse(os.path.exists(self.local_path(doc)))
		self.assertEqual(frappe.get_doc("File", doc.name).get_content(), "cloud resident")

	def test_C5_a_transient_s3_failure_is_not_a_file_not_found_error(self):
		"""The GST Return Log catches only FileNotFoundError and nulls its own pointer."""
		doc = self.make_file(file_name="c5-outage.txt", content="still exists")
		reloaded = frappe.get_doc("File", doc.name)

		for failure in (
			CloudStorageTransportError("s3 500"),
			CloudStoragePermissionError("access denied"),
		):
			with self.subTest(failure=type(failure).__name__):
				self.store.fail_with = failure
				try:
					with self.assertRaises(type(failure)) as caught:
						reloaded.get_content()
				finally:
					self.store.fail_with = None
				# Assert on what the runtime actually raised, not on the fixture: the point
				# is that `get_content` did not convert a transient failure into absence.
				self.assertNotIsInstance(caught.exception, FileNotFoundError)
				self.assertNotIsInstance(caught.exception, CloudObjectNotFound)

	def test_C5_a_genuinely_absent_object_raises_file_not_found_error(self):
		doc = self.make_file(file_name="c5-absent.txt", content="about to vanish")
		cso = self.cso_of(doc)
		self.store.objects.pop(cso.s3_key)

		with self.assertRaises(FileNotFoundError):
			frappe.get_doc("File", doc.name).get_content()

	def test_C5_the_absence_error_is_our_typed_one(self):
		doc = self.make_file(file_name="c5-typed.txt", content="vanishing")
		cso = self.cso_of(doc)
		self.store.objects.pop(cso.s3_key)

		with self.assertRaises(CloudObjectNotFound):
			frappe.get_doc("File", doc.name).get_content()


class TestC6RemoteOnlySave(ContractTestCase):
	def test_C6_saving_a_cloud_only_file_row_does_not_throw(self):
		"""`validate_file_on_disk` throws IOError in core; every re-save hits it."""
		doc = self.make_file(file_name="c6.txt", content="re-saveable")
		reloaded = frappe.get_doc("File", doc.name)
		reloaded.file_name = "c6-renamed.txt"
		reloaded.save(ignore_permissions=True)

		self.assertEqual(frappe.db.get_value("File", doc.name, "file_name"), "c6-renamed.txt")

	def test_C6_validate_file_on_disk_returns_true_without_materializing(self):
		doc = self.make_file(file_name="c6-validate.txt", content="no download please")
		reloaded = frappe.get_doc("File", doc.name)

		self.assertTrue(reloaded.validate_file_on_disk())
		self.assertEqual(materialize.tracked_paths(), {})

	def test_C6_validate_file_path_keeps_the_containment_rule(self):
		doc = self.make_file(file_name="c6-path.txt", content="contained")
		reloaded = frappe.get_doc("File", doc.name)
		reloaded.validate_file_path()

		reloaded.file_url = "/private/files/../../../etc/passwd"
		with self.assertRaises(frappe.ValidationError):
			reloaded.validate_file_path()


class TestC7MaterializedPaths(ContractTestCase):
	def _path_for(self, file_name, payload):
		doc = self.make_file(file_name=file_name, content=payload)
		return frappe.get_doc("File", doc.name).get_full_path()

	def test_C7_the_path_is_consumable_by_open(self):
		path = self._path_for("c7.txt", "openable")
		with open(path) as handle:
			self.assertEqual(handle.read(), "openable")

	def test_C7_the_path_is_consumable_by_zipfile(self):
		path = self._path_for("c7.zip", _zip_bytes())
		with zipfile.ZipFile(path) as archive:
			self.assertEqual(archive.read("inner.txt"), b"inner payload")

	def test_C7_the_path_is_consumable_by_pillow(self):
		from PIL import Image

		path = self._path_for("c7.png", _png_bytes())
		with Image.open(path) as image:
			self.assertEqual(image.size, (4, 4))

	def test_C7_the_path_is_consumable_by_openpyxl(self):
		from openpyxl import load_workbook

		path = self._path_for("c7.xlsx", _xlsx_bytes())
		self.assertEqual(load_workbook(path).active["A1"].value, "hello")

	def test_C7_the_path_is_consumable_by_csv_reader(self):
		path = self._path_for("c7.csv", "a,b\n1,2\n")
		with open(path, newline="") as handle:
			self.assertEqual(list(csv.reader(handle)), [["a", "b"], ["1", "2"]])

	def test_C7_the_cache_lives_outside_public_and_private_files(self):
		"""ADR-10 — never web-served, never tarred by the backup generator."""
		path = self._path_for("c7-location.txt", "somewhere safe")
		self.assertIn("cloud_storage_cache", path)
		self.assertNotIn(os.path.join("public", "files"), path)
		self.assertNotIn(os.path.join("private", "files"), path)


class TestC8WriteThrough(ContractTestCase):
	def test_C8_a_write_against_the_returned_path_is_flushed_back(self):
		"""`bank_statement_import.py:202-212` writes straight into `get_full_path()`."""
		doc = self.make_file(file_name="c8.txt", content="before")
		reloaded = frappe.get_doc("File", doc.name)
		original_cso = reloaded.cloud_storage_object

		path = reloaded.get_full_path()
		with open(path, "w") as handle:
			handle.write("after")

		writeback.flush_request()

		row = frappe.db.get_value("File", doc.name, ["cloud_storage_object", "content_hash"], as_dict=True)
		self.assertNotEqual(row.cloud_storage_object, original_cso)
		self.assertEqual(row.content_hash, get_content_hash(b"after"))
		self.assertEqual(frappe.get_doc("File", doc.name).get_content(), "after")

	def test_C8_an_unchanged_path_does_not_create_a_new_object(self):
		doc = self.make_file(file_name="c8-unchanged.txt", content="stable")
		reloaded = frappe.get_doc("File", doc.name)
		original_cso = reloaded.cloud_storage_object

		reloaded.get_full_path()
		writeback.flush_request()

		self.assertEqual(frappe.db.get_value("File", doc.name, "cloud_storage_object"), original_cso)


class TestC9CanonicalUrls(ContractTestCase):
	def test_C9_every_written_url_is_canonical(self):
		"""`attach_files_to_document`, `set_is_private` and `find_file_by_url` all depend on it."""
		for is_private in (0, 1):
			with self.subTest(is_private=is_private):
				doc = self.make_file(
					file_name=f"c9-{is_private}.txt", content=f"canonical {is_private}", is_private=is_private
				)
				expected_prefix = "/private/files/" if is_private else "/files/"
				stored = frappe.db.get_value("File", doc.name, "file_url")
				self.assertTrue(stored.startswith(expected_prefix), stored)

	def test_C9_no_signed_or_endpoint_url_is_ever_stored(self):
		doc = self.make_file(file_name="c9-signed.txt", content="never a signed url")
		stored = frappe.db.get_value("File", doc.name, "file_url")

		for forbidden in ("http://", "https://", "/api/method/", "X-Amz-", "?"):
			with self.subTest(forbidden=forbidden):
				self.assertNotIn(forbidden, stored)

	def test_C9_is_private_is_derivable_from_the_url_as_core_expects(self):
		"""`file.py:784 set_is_private` derives privacy from `startswith('/private')`."""
		for is_private in (0, 1):
			with self.subTest(is_private=is_private):
				doc = self.make_file(
					file_name=f"c9-derive-{is_private}.txt",
					content=f"derive {is_private}",
					is_private=is_private,
				)
				stored = frappe.db.get_value("File", doc.name, "file_url")
				self.assertEqual(int(stored.startswith("/private")), is_private)


class TestC10PermissionGate(ContractTestCase):
	"""C10 — no URL is issued before the File's own read gate has run.

	P2 ships one URL-issuing surface, the legacy compat endpoint; the full
	`download_private_file` interception with its exactly-one access log is P3's spike and
	is tested there. What is asserted here is the invariant P2 must not break: the gate
	runs first, and nothing in the runtime core hands out an object URL without it.
	"""

	def test_C10_the_legacy_endpoint_refuses_a_user_without_read_permission(self):
		from cloud_file_storage.api import compat

		doc = self.make_file(file_name="c10-secret.txt", content="classified")
		cso = self.cso_of(doc)

		user = "cfs_c10_probe@test.local"
		if not frappe.db.exists("User", user):
			probe = frappe.new_doc("User")
			probe.email = user
			probe.first_name = "C10 Probe"
			probe.send_welcome_email = 0
			probe.insert(ignore_permissions=True)
			probe.add_roles("Desk User")
		self.addCleanup(frappe.set_user, "Administrator")

		frappe.set_user(user)
		with self.assertRaises(frappe.PermissionError):
			compat.legacy_generate_file(key=cso.s3_key, file_name="c10-secret.txt")

	def test_C10_an_authorised_caller_gets_a_short_lived_redirect(self):
		from cloud_file_storage.api import compat

		doc = self.make_file(file_name="c10-mine.txt", content="mine")
		cso = self.cso_of(doc)

		frappe.local.response = frappe._dict()
		compat.legacy_generate_file(key=cso.s3_key, file_name="c10-mine.txt")

		self.assertEqual(frappe.local.response["type"], "redirect")
		# The TTL the runtime passed, not a default the fake invented.
		issued = self.store.presign_calls[-1]
		self.assertEqual(issued["ttl"], 300)
		# The redirect is the URL that was minted for THIS object, not a rewrite of it.
		# (`X-Amz-Expires=300` here would only echo the ttl the fake was handed.)
		self.assertEqual(frappe.local.response["location"], issued["url"])
		self.assertEqual(issued["key"], cso.s3_key)

	def test_C10_the_runtime_core_issues_no_url_outside_that_gate(self):
		import ast
		from pathlib import Path

		import cloud_file_storage

		package = Path(cloud_file_storage.__file__).resolve().parent
		callers = []
		for path in package.rglob("*.py"):
			# Exclude by DIRECTORY, not by basename substring. `"test" in path.name` misses
			# every helper under tests/ whose name lacks it — utils.py, backup_utils.py,
			# desk_utils.py, permission_utils.py, rehearsal.py. None trips the matcher
			# today, but the day a fake legitimately calls generate_presigned_url this
			# contract fails naming a test helper, and the tempting fix is to add that
			# helper to SANCTIONED_PRESIGN_CALLERS — which is how an allow-list stops
			# meaning what its name says.
			rel_path = path.relative_to(package)
			relative = rel_path.as_posix()
			if "tests" in rel_path.parts or relative == "storage/engine.py":
				continue
			tree = ast.parse(path.read_text(encoding="utf-8"))
			for node in ast.walk(tree):
				# `engine.presign_get(...)` — the app's own wrapper — and
				# `client.generate_presigned_url(...)` — boto3 directly. This gate was written
				# when `presign_get` was the only wrapper in existence, so "no module mentions
				# presign_get" and "no module mints a URL" were the same property.
				# `backup/restore.py` broke that equivalence and the test stayed green (H-1).
				if isinstance(node, ast.Attribute) and _is_minting_name(node.attr):
					callers.append(relative)
				# The import form the attribute walk misses. BOTH names, symmetrically:
				# `presign_get` was matched in three syntactic forms and `generate_presigned*`
				# in one, so `from botocore.signers import generate_presigned_url` followed by
				# a bare call escaped entirely — H-1's shape exactly, a function that was
				# always reachable by a route the matcher did not know (M-8).
				if isinstance(node, ast.ImportFrom):
					if any(_is_minting_name(alias.name) for alias in node.names):
						callers.append(relative)
				# a bare `presign_get(...)` / `generate_presigned_url(...)` call, however it
				# got into scope
				if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
					if _is_minting_name(node.func.id):
						callers.append(relative)

		self.assertEqual(
			sorted(set(callers)),
			SANCTIONED_PRESIGN_CALLERS,
			"a new presigned-URL caller must come with its own permission-gate test",
		)

	@refusal_guard
	def test_C10_the_matcher_catches_every_minting_form(self):
		"""Symmetry, driven against synthetic source (M-8).

		`presign_get` was matched as an attribute, as a `from … import`, and as a bare call;
		`generate_presigned*` only as an attribute. So
		`from botocore.signers import generate_presigned_url` followed by a bare call escaped
		the contract completely — H-1's shape a second time, a reachable route the matcher did
		not know. Each form is asserted for BOTH names, because the asymmetry was the defect
		and a test covering only the form that escaped would let the mirror image back in.
		"""
		import ast

		forms = {
			"attribute": "def f(c):\n\treturn c.{name}('get_object')\n",
			"import-then-bare-call": "from x import {name}\n\ndef f(c):\n\treturn {name}(c)\n",
			"bare call only": "def f(c, {name}):\n\treturn {name}(c)\n",
		}
		names = ["presign_get", "generate_presigned_url", "generate_presigned_post"]

		for name in names:
			for label, template in forms.items():
				with self.subTest(name=name, form=label):
					source = template.format(name=name)
					self.assertTrue(
						self._module_mints(ast.parse(source)),
						f"the C10 walk misses {name} in the {label} form",
					)

		# And the control: a module that mints nothing is not flagged, or every assertion
		# above would hold for any input at all.
		self.assertFalse(self._module_mints(ast.parse("def f(c):\n\treturn c.head_object()\n")))

	@staticmethod
	def _module_mints(tree) -> bool:
		"""The C10 walk's three branches, over one parsed module.

		Kept in step with the scan in `test_C10_the_runtime_core_issues_no_url_outside_that_gate`
		by sharing `_is_minting_name`, which is the part that was asymmetric.
		"""
		import ast

		for node in ast.walk(tree):
			if isinstance(node, ast.Attribute) and _is_minting_name(node.attr):
				return True
			if isinstance(node, ast.ImportFrom) and any(_is_minting_name(alias.name) for alias in node.names):
				return True
			if (
				isinstance(node, ast.Call)
				and isinstance(node.func, ast.Name)
				and _is_minting_name(node.func.id)
			):
				return True
		return False

	def test_C10_the_public_renderer_refuses_a_private_row(self):
		"""A19 — the public path's gate. It has no permission check by design, so the
		`is_private = 0` requirement IS the security boundary and it is asserted here
		alongside the caller allow-list above."""
		from cloud_file_storage.serving.public import PublicFileRenderer
		from cloud_file_storage.tests.request_utils import http_request

		private_doc = self.make_file(file_name="c10-private.txt", content="not public", is_private=1)
		public_url = private_doc.file_url.replace("/private/files/", "/files/")

		with http_request(public_url):
			renderer = PublicFileRenderer(public_url.lstrip("/"))
			self.assertFalse(renderer.can_render())
		self.assertEqual(self.store.presign_calls, [])


class TestC11OverwriteInPlace(ContractTestCase):
	def test_C11_overwrite_keeps_the_url_and_repoints_the_bytes(self):
		"""`gst_return_log.py:103` then stores `file.file_url` in its own field."""
		doc = self.make_file(file_name="c11.json", content='{"v": 1}')
		original_url = doc.file_url
		original_cso = doc.cloud_storage_object

		reloaded = frappe.get_doc("File", doc.name)
		reloaded.save_file(content='{"v": 2}', overwrite=True)

		self.assertEqual(reloaded.file_url, original_url)
		self.assertNotEqual(frappe.db.get_value("File", doc.name, "cloud_storage_object"), original_cso)
		self.assertEqual(frappe.get_doc("File", doc.name).get_content(), '{"v": 2}')

	def test_C11_the_previous_object_survives_until_gc(self):
		"""Refcounting tolerates an object whose content changed under a fixed URL."""
		doc = self.make_file(file_name="c11-survive.json", content='{"v": 1}')
		original_cso = doc.cloud_storage_object

		frappe.get_doc("File", doc.name).save_file(content='{"v": 2}', overwrite=True)

		self.assertEqual(frappe.db.get_value(CSO_DOCTYPE, original_cso, "status"), "pending_delete")
		self.assertTrue(self.store.contains(frappe.get_doc(CSO_DOCTYPE, original_cso)))


class TestC12BothWriteHooks(ContractTestCase):
	def test_C12_the_document_convention_is_implemented(self):
		doc = self.make_file(file_name="c12-doc.txt", content="document convention")
		self.assertTrue(frappe.db.get_value("File", doc.name, "cloud_storage_object"))

	def test_C12_the_file_manager_convention_returns_the_expected_dict(self):
		from cloud_file_storage import core_hooks

		result = core_hooks.write_file(
			"c12-legacy.json", b'{"irn": 1}', content_type="application/json", is_private=1
		)
		self.assertEqual(set(result), {"file_name", "file_url"})
		self.assertTrue(result["file_url"].startswith("/private/files/"))

	def test_C12_the_e_waybill_shape_produces_a_linked_readable_file(self):
		from frappe.utils.file_manager import save_file

		doc = save_file("c12-ewaybill.json", b'{"irn": "abc"}', None, None, is_private=1)
		self.track(doc)

		self.assertTrue(frappe.db.get_value("File", doc.name, "cloud_storage_object"))
		self.assertEqual(frappe.get_doc("File", doc.name).get_content(), '{"irn": "abc"}')


class TestC13RefcountSafeDeletes(ContractTestCase):
	def test_C13_deleting_one_of_two_sharers_leaves_the_object_alone(self):
		first = self.make_file(file_name="c13-a.txt", content="shared bytes")
		second = self.make_file(file_name="c13-b.txt", content="shared bytes", ignore_duplicate_entry_error=1)
		cso_name = frappe.db.get_value("File", first.name, "cloud_storage_object")
		self.assertEqual(frappe.db.get_value("File", second.name, "cloud_storage_object"), cso_name)

		frappe.delete_doc("File", first.name, ignore_permissions=True, force=True)

		self.assertEqual(frappe.db.get_value(CSO_DOCTYPE, cso_name, "status"), "uploaded")
		self.assertTrue(self.store.contains(frappe.get_doc(CSO_DOCTYPE, cso_name)))

	def test_C13_deleting_the_last_reference_schedules_but_never_deletes_inline(self):
		doc = self.make_file(file_name="c13-last.txt", content="last reference")
		cso_name = doc.cloud_storage_object

		frappe.delete_doc("File", doc.name, ignore_permissions=True, force=True)

		self.assertEqual(frappe.db.get_value(CSO_DOCTYPE, cso_name, "status"), "pending_delete")
		self.assertEqual(self.store.deleted, [], "the delete hook must never touch the bucket")
		self.assertTrue(self.store.contains(frappe.get_doc(CSO_DOCTYPE, cso_name)))

	def test_C13_the_hook_is_authoritative_and_registered(self):
		from frappe.utils import get_hook_method

		from cloud_file_storage import core_hooks

		self.assertIs(get_hook_method("delete_file_data_content"), core_hooks.delete_file_data_content)


class TestC14RemoteExistence(ContractTestCase):
	def test_C14_exists_on_disk_reports_remote_existence(self):
		"""`save_file` uses this to decide whether to reuse a duplicate's `file_url`."""
		doc = self.make_file(file_name="c14.txt", content="remote but present")
		reloaded = frappe.get_doc("File", doc.name)

		self.assertFalse(os.path.exists(self.local_path(reloaded)))
		self.assertTrue(reloaded.exists_on_disk())

	def test_C14_it_reports_false_for_an_object_that_was_never_uploaded(self):
		doc = self.make_file(file_name="c14-pending.txt", content="not uploaded")
		cso_name = doc.cloud_storage_object
		frappe.db.set_value(CSO_DOCTYPE, cso_name, "status", "pending_upload")

		self.assertFalse(frappe.get_doc("File", doc.name).exists_on_disk())

	def test_C14_dedup_reuses_the_url_of_a_cloud_only_twin(self):
		first = self.make_file(file_name="c14-first.txt", content="twin bytes")
		second = frappe.get_doc(
			{
				"doctype": "File",
				"file_name": "c14-second.txt",
				"content": "twin bytes",
				"is_private": 1,
				"ignore_duplicate_entry_error": 1,
			}
		).insert(ignore_permissions=True)
		self.track(second)

		self.assertEqual(second.file_url, first.file_url)

	def test_C14_it_does_not_hit_the_network(self):
		doc = self.make_file(file_name="c14-nonetwork.txt", content="no head please")
		reloaded = frappe.get_doc("File", doc.name)
		self.store.fail_with = CloudStorageTransportError("no network calls allowed here")
		try:
			self.assertTrue(reloaded.exists_on_disk())
		finally:
			self.store.fail_with = None


class TestC15ContentHash(ContractTestCase):
	def test_C15_content_hash_is_the_md5_of_the_bytes(self):
		payload = _png_bytes()
		doc = self.make_file(file_name="c15.png", content=payload)

		self.assertEqual(frappe.db.get_value("File", doc.name, "content_hash"), get_content_hash(payload))

	def test_C15_the_object_dedup_index_uses_the_same_value(self):
		payload = b"c15 dedup identity"
		doc = self.make_file(file_name="c15-dedup.bin", content=payload)
		cso = self.cso_of(doc)

		self.assertEqual(cso.content_hash_md5, get_content_hash(payload))

	def test_C15_generate_content_hash_does_not_throw_for_a_cloud_only_row(self):
		doc = self.make_file(file_name="c15-generate.txt", content="no local file here")
		reloaded = frappe.get_doc("File", doc.name)
		reloaded.content_hash = None
		reloaded.generate_content_hash()

		self.assertEqual(reloaded.content_hash, get_content_hash(b"no local file here"))


class TestC16PrivacyFlip(ContractTestCase):
	def test_C16_flipping_is_private_works_without_shutil_move(self):
		doc = self.make_file(file_name="c16.txt", content="flip me", is_private=1)
		old_cso = doc.cloud_storage_object

		reloaded = frappe.get_doc("File", doc.name)
		reloaded.is_private = 0
		reloaded.save(ignore_permissions=True)

		self.assertTrue(reloaded.file_url.startswith("/files/"))
		new_cso_name = frappe.db.get_value("File", doc.name, "cloud_storage_object")
		self.assertNotEqual(new_cso_name, old_cso)
		new_cso = frappe.get_doc(CSO_DOCTYPE, new_cso_name)
		self.assertEqual(new_cso.visibility, "public")
		self.assertIn("/pub/", new_cso.s3_key)
		self.assertTrue(self.store.contains(new_cso))

	def test_C16_the_flipped_file_is_still_readable(self):
		doc = self.make_file(file_name="c16-readable.txt", content="still readable")
		reloaded = frappe.get_doc("File", doc.name)
		reloaded.is_private = 0
		reloaded.save(ignore_permissions=True)

		self.assertEqual(frappe.get_doc("File", doc.name).get_content(), "still readable")

	def test_C16_every_content_hash_sibling_is_relinked_in_the_same_transaction(self):
		"""A16 — core rewrites those rows' URLs, so their object must follow."""
		doc = self.make_file(file_name="c16-source.txt", content="sibling flip")
		sibling = self.make_file(
			file_name="c16-sibling.txt", content="sibling flip", ignore_duplicate_entry_error=1
		)
		old_cso = frappe.db.get_value("File", doc.name, "cloud_storage_object")
		self.assertEqual(frappe.db.get_value("File", sibling.name, "cloud_storage_object"), old_cso)

		reloaded = frappe.get_doc("File", doc.name)
		reloaded.is_private = 0
		reloaded.save(ignore_permissions=True)

		new_cso = frappe.db.get_value("File", doc.name, "cloud_storage_object")
		self.assertNotEqual(new_cso, old_cso)
		self.assertEqual(frappe.db.get_value("File", sibling.name, "cloud_storage_object"), new_cso)

	def test_C16_the_old_object_is_only_released_through_the_refcount_path(self):
		doc = self.make_file(file_name="c16-release.txt", content="release path")
		old_cso = doc.cloud_storage_object

		reloaded = frappe.get_doc("File", doc.name)
		reloaded.is_private = 0
		reloaded.save(ignore_permissions=True)

		status = frappe.db.get_value(CSO_DOCTYPE, old_cso, "status")
		self.assertEqual(status, "pending_delete")
		self.assertEqual(self.store.deleted, [], "a privacy flip must never delete bytes inline")


class TestC17WritablePathFallthrough(ContractTestCase):
	"""C17 — the erpnext `stock_ledger.py:416` old-app sentinel no longer fires.

	The contract offers two ways out: keep the sentinel working, or make
	`get_full_path()` + write S3-aware so the fallthrough at `:420-422` is safe. This is
	the second, which is the one the approved architecture chose.
	"""

	def test_C17_the_old_app_name_sentinel_does_not_appear_in_any_url(self):
		doc = self.make_file(file_name="c17.txt", content="no sentinel")
		self.assertNotIn("/frappe_s3_attachment.", frappe.db.get_value("File", doc.name, "file_url"))

	def test_C17_the_stock_ledger_fallthrough_shape_is_safe(self):
		"""`open(get_full_path(), "wb")` on a gz payload, inside one job."""
		doc = self.make_file(file_name="c17-sle.json.gz", content=gzip.compress(b'{"sle": 1}'))
		reloaded = frappe.get_doc("File", doc.name)
		original_cso = reloaded.cloud_storage_object

		path = reloaded.get_full_path()
		with open(path, "wb") as handle:
			handle.write(gzip.compress(b'{"sle": 2}'))

		writeback.flush_job()

		self.assertNotEqual(frappe.db.get_value("File", doc.name, "cloud_storage_object"), original_cso)
		self.assertEqual(
			gzip.decompress(frappe.get_doc("File", doc.name).get_content(encodings=[])), b'{"sle": 2}'
		)


class TestC18BusinessFieldPointers(ContractTestCase):
	"""C18 — `stock_ledger.py:308-318` looks Files up by `file_url` + `attached_to_field`."""

	def test_C18_a_stored_url_still_finds_the_row_after_an_overwrite(self):
		doc = self.make_file(file_name="c18.json", content='{"v": 1}')
		stored_pointer = doc.file_url

		frappe.get_doc("File", doc.name).save_file(content='{"v": 2}', overwrite=True)

		found = frappe.db.get_value("File", {"file_url": stored_pointer}, "name")
		self.assertEqual(found, doc.name)
		self.assertEqual(frappe.get_doc("File", found).get_content(), '{"v": 2}')

	def test_C18_a_privacy_flip_propagates_the_new_url_to_the_parent_field(self):
		parent = frappe.get_doc(
			{"doctype": "ToDo", "description": "cfs c18 parent", "status": "Open"}
		).insert(ignore_permissions=True)
		self.addCleanup(lambda: frappe.db.delete("ToDo", {"name": parent.name}))

		doc = frappe.get_doc(
			{
				"doctype": "File",
				"file_name": "c18-attach.txt",
				"content": "attached bytes",
				"is_private": 1,
				"attached_to_doctype": "ToDo",
				"attached_to_name": parent.name,
				"attached_to_field": "description",
			}
		).insert(ignore_permissions=True)
		self.track(doc)

		reloaded = frappe.get_doc("File", doc.name)
		reloaded.is_private = 0
		reloaded.save(ignore_permissions=True)

		self.assertEqual(frappe.db.get_value("ToDo", parent.name, "description"), reloaded.file_url)

	def test_C18_url_siblings_are_relinked_so_one_url_means_one_object(self):
		doc = self.make_file(file_name="c18-sibling.json", content='{"v": 1}')
		original_cso = doc.cloud_storage_object

		sibling = frappe.get_doc(
			{
				"doctype": "File",
				"file_name": "c18-sibling.json",
				"file_url": doc.file_url,
				"is_private": 1,
				"content_hash": doc.content_hash,
			}
		)
		sibling.flags.ignore_duplicate_entry_error = True
		sibling.insert(ignore_permissions=True)
		self.track(sibling)
		objects.adopt_references(original_cso, [sibling.name])

		frappe.get_doc("File", doc.name).save_file(content='{"v": 2}', overwrite=True)

		new_cso = frappe.db.get_value("File", doc.name, "cloud_storage_object")
		self.assertEqual(frappe.db.get_value("File", sibling.name, "cloud_storage_object"), new_cso)


class TestC19AppSourceFixturesUntouched(ContractTestCase):
	"""C19 — `frappe.get_file_json` on app-source paths is explicitly out of scope.

	India Compliance reads its fixtures straight off disk from the app package. Those are
	not File records and must keep working byte-for-byte; this app must not patch that
	function or route it through storage.
	"""

	def test_C19_get_file_json_still_reads_app_source_paths_from_disk(self):
		path = frappe.get_app_path("frappe", "core", "doctype", "file", "file.json")
		self.assertEqual(frappe.get_file_json(path)["name"], "File")

	def test_C19_the_app_does_not_patch_get_file_json(self):
		from pathlib import Path

		import cloud_file_storage

		package = Path(cloud_file_storage.__file__).resolve().parent
		offenders = [
			source.relative_to(package).as_posix()
			for source in package.rglob("*.py")
			if "test" not in source.name and patches_get_file_json(source.read_text(encoding="utf-8"))
		]
		self.assertEqual(sorted(set(offenders)), [])

	def test_C19_the_patch_detector_still_catches_every_patching_form(self):
		"""The detector ignores prose, so prove it still bites on every real patch.

		Without this, the docstring exclusion could quietly become the reason
		`test_C19_the_app_does_not_patch_get_file_json` is green.
		"""
		self.assertTrue(patches_get_file_json('patch("frappe.get_file_json")'))
		self.assertTrue(patches_get_file_json("frappe.get_file_json = ours"))
		self.assertTrue(patches_get_file_json("from frappe import get_file_json"))
		self.assertTrue(patches_get_file_json('frappe.get_attr("frappe.get_file_json")'))
		self.assertTrue(patches_get_file_json('def f():\n\t"""doc"""\n\tx = "frappe.get_file_json"'))

		self.assertFalse(patches_get_file_json('"""We never touch frappe.get_file_json."""'))
		self.assertFalse(patches_get_file_json("# frappe.get_file_json stays core's\npass"))
		self.assertFalse(patches_get_file_json('def f():\n\t"""Not get_file_json either."""'))

	def test_C19_reading_an_app_source_path_never_touches_the_object_store(self):
		set_mode("S3_ONLY")
		self.store.fail_with = CloudStorageTransportError("no storage call expected")
		try:
			path = frappe.get_app_path("frappe", "core", "doctype", "file", "file.json")
			self.assertTrue(frappe.get_file_json(path))
		finally:
			self.store.fail_with = None
